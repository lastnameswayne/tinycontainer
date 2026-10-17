import io
from openpyxl import Workbook, load_workbook

# first output marks "imports done"
print("ready: d-openpyxl")

wb = Workbook()
ws = wb.active
for i in range(20_000):
    ws.append([i, i * 1.5, f"row{i}"])
buf = io.BytesIO()
wb.save(buf)
back = load_workbook(io.BytesIO(buf.getvalue()), read_only=True).active
print(f"xlsx bytes={len(buf.getvalue())} rows={sum(1 for _ in back.iter_rows())}")
